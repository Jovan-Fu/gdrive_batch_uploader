import os
import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional, Callable, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.google_client import GoogleCloudClient, extract_google_id


class HierarchyReportExporter:
    """
    负责将上传结果按【总文件夹 -> 分类子文件夹 -> 单个文件】的层级结构：
    1. 导出到本地 Excel (.xlsx) 表格
    2. 写入/追加到云端 Google Sheets (谷歌表格)
    包含两个工作表：
    - 「层级分栏总表」：从左到右按【总文件夹 -> 子文件夹 -> 单个文件】三级分栏，支持同级单元格纵向合并
    - 「树状目录层级表」：从上到下按 📁总文件夹 -> 📂子文件夹 -> 📄单文件 树状缩进排列
    """

    def __init__(self, log_callback: Optional[Callable[[str, str], None]] = None):
        self.log_callback = log_callback

    def _log(self, msg: str, level: str = "INFO"):
        if self.log_callback:
            self.log_callback(msg, level)

    def export_local_excel(
        self,
        batch_result: Dict[str, Any],
        output_path: str,
        merge_hierarchy_cells: bool = True,
    ) -> str:
        """
        生成本地层级排版的 Excel 文件 (.xlsx)
        batch_result 结构:
        {
            "batch_time": "2026-10-27 09:30:00",
            "master_folder": {"id": "...", "name": "...", "webViewLink": "..."},
            "planned_folders": M,
            "planned_files_per_folder": N,
            "total_uploaded_files": X,
            "summary_note": "...",
            "subfolders": [
                {
                    "index": 1,
                    "id": "...",
                    "name": "001",
                    "webViewLink": "...",
                    "files": [
                        {
                            "index_in_folder": 1,
                            "global_index": 1,
                            "name": "a.mp4",
                            "id": "...",
                            "downloadLink": "...",
                            "webViewLink": "...",
                            "size_str": "1.2 MB",
                            "permission_status": "已开启公开下载权限",
                            "upload_time": "2026-10-27 09:30:05"
                        }
                    ]
                }
            ]
        }
        """
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        # Sheet 1: 层级分栏总表
        ws_col = wb.active
        ws_col.title = "层级分栏总表"
        self._build_excel_column_hierarchy_sheet(
            ws_col, batch_result, merge_hierarchy_cells
        )

        # Sheet 2: 树状目录层级表
        ws_tree = wb.create_sheet(title="树状目录层级表")
        self._build_excel_tree_hierarchy_sheet(ws_tree, batch_result)

        wb.save(output_path)
        self._log(f"本地层级 Excel 表格已保存：{output_path}", "SUCCESS")
        return output_path

    def _build_excel_column_hierarchy_sheet(
        self,
        ws,
        batch_result: Dict[str, Any],
        merge_hierarchy_cells: bool = True,
    ):
        """构建左->右三级分栏层级表（总文件夹 -> 子文件夹 -> 单文件）"""
        thin_side = Side(border_style="thin", color="D0D7DE")
        border_all = Border(
            left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
        )

        # 颜色主题
        fill_l1_hdr = PatternFill("solid", fgColor="1B4F72")  # 深海蓝 - 总文件夹一级头
        fill_l2_hdr = PatternFill("solid", fgColor="117864")  # 深青绿 - 子文件夹一级头
        fill_l3_hdr = PatternFill("solid", fgColor="2E4053")  # 深灰蓝 - 单个文件一级头

        fill_sub_hdr = PatternFill("solid", fgColor="EBF5FB")
        fill_l1_cell = PatternFill("solid", fgColor="F4F9FD")
        fill_l2_cell_even = PatternFill("solid", fgColor="E8F8F5")
        fill_l2_cell_odd = PatternFill("solid", fgColor="FEF9E7")

        font_group_hdr = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        font_col_hdr = Font(name="Microsoft YaHei", size=10, bold=True, color="1C2833")
        font_bold = Font(name="Microsoft YaHei", size=10, bold=True, color="1B4F72")
        font_regular = Font(name="Microsoft YaHei", size=10, color="2C3E50")
        font_link = Font(
            name="Microsoft YaHei", size=9, color="0563C1", underline="single"
        )

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # 第1行：一级大类层级分组表头
        ws.merge_cells("A1:C1")
        ws["A1"] = "【第一层级】云端总文件夹 (Master Folder)"
        ws.merge_cells("D1:G1")
        ws["D1"] = "【第二层级】分类子文件夹 (Sub-Folder)"
        ws.merge_cells("H1:N1")
        ws["H1"] = "【第三层级】夹内单个文件 (Individual File & 下载链接)"

        for col_idx in range(1, 15):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = font_group_hdr
            cell.alignment = align_center
            cell.border = border_all
            if col_idx <= 3:
                cell.fill = fill_l1_hdr
            elif col_idx <= 7:
                cell.fill = fill_l2_hdr
            else:
                cell.fill = fill_l3_hdr
        ws.row_dimensions[1].height = 28

        # 第2行：二级具体字段表头
        headers = [
            "总文件夹名称",
            "总文件夹 ID",
            "总文件夹链接",
            "子文件夹序号",
            "子文件夹名称",
            "子文件夹 ID",
            "子文件夹链接",
            "夹内序号",
            "单个文件名称",
            "单个文件 ID",
            "单文件直接下载链接 (Download直链)",
            "单文件在线预览链接 (View链接)",
            "文件大小",
            "权限状态 / 上传时间",
        ]
        for col_idx, h_text in enumerate(headers, start=1):
            c = ws.cell(row=2, column=col_idx, value=h_text)
            c.font = font_col_hdr
            c.fill = fill_sub_hdr
            c.alignment = align_center
            c.border = border_all
        ws.row_dimensions[2].height = 25

        master = batch_result.get("master_folder", {})
        master_name = master.get("name", "")
        master_id = master.get("id", "")
        master_link = master.get("webViewLink", "")

        subfolders = batch_result.get("subfolders", [])
        current_row = 3
        batch_start_row = current_row

        for sf_idx, sf in enumerate(subfolders):
            sf_order = sf.get("index", sf_idx + 1)
            sf_name = sf.get("name", "")
            sf_id = sf.get("id", "")
            sf_link = sf.get("webViewLink", "")
            files = sf.get("files", [])

            sf_start_row = current_row
            sf_fill = fill_l2_cell_even if sf_idx % 2 == 0 else fill_l2_cell_odd

            if not files:
                # 空文件夹行（如果存在）
                row_vals = [
                    master_name,
                    master_id,
                    master_link,
                    sf_order,
                    sf_name,
                    sf_id,
                    sf_link,
                    "-",
                    "(无文件)",
                    "-",
                    "-",
                    "-",
                    "-",
                    "无文件",
                ]
                for col_idx, val in enumerate(row_vals, start=1):
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    cell.font = font_regular
                    cell.border = border_all
                    cell.alignment = align_center if col_idx in (4, 8, 13) else align_left
                ws.row_dimensions[current_row].height = 22
                current_row += 1
                continue

            for f_item in files:
                row_vals = [
                    master_name,
                    master_id,
                    master_link,
                    sf_order,
                    sf_name,
                    sf_id,
                    sf_link,
                    f_item.get("index_in_folder", 1),
                    f_item.get("name", ""),
                    f_item.get("id", ""),
                    f_item.get("downloadLink", ""),
                    f_item.get("webViewLink", ""),
                    f_item.get("size_str", ""),
                    f"{f_item.get('permission_status', '已开启公开权限')} ({f_item.get('upload_time', '')})",
                ]
                for col_idx, val in enumerate(row_vals, start=1):
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    cell.border = border_all

                    if col_idx in (3, 7, 11, 12) and isinstance(val, str) and val.startswith("http"):
                        cell.hyperlink = val
                        cell.font = font_link
                    elif col_idx in (1, 5, 9):
                        cell.font = font_bold if col_idx in (1, 5) else font_regular
                    else:
                        cell.font = font_regular

                    if col_idx <= 3:
                        cell.fill = fill_l1_cell
                        cell.alignment = align_center if col_idx == 1 else align_left
                    elif col_idx <= 7:
                        cell.fill = sf_fill
                        cell.alignment = align_center if col_idx in (4, 5) else align_left
                    else:
                        cell.alignment = align_center if col_idx in (8, 13) else align_left

                ws.row_dimensions[current_row].height = 22
                current_row += 1

            sf_end_row = current_row - 1
            if merge_hierarchy_cells and sf_end_row > sf_start_row:
                for col_idx in (4, 5, 6, 7):
                    ws.merge_cells(
                        start_row=sf_start_row,
                        start_column=col_idx,
                        end_row=sf_end_row,
                        end_column=col_idx,
                    )

        batch_end_row = current_row - 1
        if merge_hierarchy_cells and batch_end_row > batch_start_row:
            for col_idx in (1, 2, 3):
                ws.merge_cells(
                    start_row=batch_start_row,
                    start_column=col_idx,
                    end_row=batch_end_row,
                    end_column=col_idx,
                )

        # 列宽设置
        col_widths = [22, 26, 38, 12, 18, 26, 38, 10, 28, 28, 46, 42, 12, 30]
        for idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = "A3"

    def _build_excel_tree_hierarchy_sheet(self, ws, batch_result: Dict[str, Any]):
        """构建上->下树状缩进目录表（📁总文件夹 -> 📂子文件夹 -> 📄单文件）"""
        thin_side = Side(border_style="thin", color="D0D7DE")
        border_all = Border(
            left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
        )

        fill_hdr = PatternFill("solid", fgColor="1C2833")
        fill_l1 = PatternFill("solid", fgColor="D6EAF8")  # 总文件夹行：柔和蓝
        fill_l2 = PatternFill("solid", fgColor="D1F2EB")  # 子文件夹行：柔和绿
        fill_l3 = PatternFill("solid", fgColor="FFFFFF")  # 单个文件行：纯白

        font_hdr = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        font_l1 = Font(name="Microsoft YaHei", size=10, bold=True, color="154360")
        font_l2 = Font(name="Microsoft YaHei", size=10, bold=True, color="0E6251")
        font_l3 = Font(name="Microsoft YaHei", size=10, color="2C3E50")
        font_link = Font(
            name="Microsoft YaHei", size=9, color="0563C1", underline="single"
        )

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        headers = [
            "层级级别",
            "层级树状目录展示",
            "节点名称 (文件夹名/文件名)",
            "云端 ID (Folder / File ID)",
            "单文件直接下载链接 (Download直链)",
            "云端打开/预览链接 (总夹/子夹/单文件链接)",
            "所属上级路径",
            "文件大小 / 数量统计",
            "权限与时间说明",
        ]
        for col_idx, h_text in enumerate(headers, start=1):
            c = ws.cell(row=1, column=col_idx, value=h_text)
            c.font = font_hdr
            c.fill = fill_hdr
            c.alignment = align_center
            c.border = border_all
        ws.row_dimensions[1].height = 26

        master = batch_result.get("master_folder", {})
        master_name = master.get("name", "")
        master_id = master.get("id", "")
        master_link = master.get("webViewLink", "")
        subfolders = batch_result.get("subfolders", [])
        total_files = batch_result.get("total_uploaded_files", 0)

        rows_data = []
        # 1. 总文件夹行
        rows_data.append(
            (
                1,
                [
                    "第1级：总文件夹",
                    f"📁 [总文件夹] {master_name}",
                    master_name,
                    master_id,
                    "(总文件夹请使用右侧打开链接)",
                    master_link,
                    "云端根目录",
                    f"共 {len(subfolders)} 个子文件夹 / {total_files} 个文件",
                    f"已开启公开访问权限 | {batch_result.get('batch_time', '')}",
                ],
            )
        )

        # 2. 遍历每个子文件夹及其下属文件
        for sf in subfolders:
            sf_idx = sf.get("index", 1)
            sf_name = sf.get("name", "")
            sf_id = sf.get("id", "")
            sf_link = sf.get("webViewLink", "")
            files = sf.get("files", [])

            rows_data.append(
                (
                    2,
                    [
                        f"第2级：子文件夹 #{sf_idx}",
                        f"    📂 [{sf_idx:02d}] {sf_name}",
                        sf_name,
                        sf_id,
                        "(子文件夹请使用右侧打开链接)",
                        sf_link,
                        master_name,
                        f"本夹内共 {len(files)} 个文件",
                        "已开启公开访问权限",
                    ],
                )
            )

            for f_item in files:
                f_idx = f_item.get("index_in_folder", 1)
                f_name = f_item.get("name", "")
                f_id = f_item.get("id", "")
                dl_link = f_item.get("downloadLink", "")
                vw_link = f_item.get("webViewLink", "")
                rows_data.append(
                    (
                        3,
                        [
                            f"第3级：单个文件 ({f_idx}/{len(files)})",
                            f"        📄 ({f_idx}) {f_name}",
                            f_name,
                            f_id,
                            dl_link,
                            vw_link,
                            f"{master_name} / {sf_name}",
                            f_item.get("size_str", ""),
                            f"{f_item.get('permission_status', '')} | {f_item.get('upload_time', '')}",
                        ],
                    )
                )

        for r_offset, (level, r_vals) in enumerate(rows_data, start=2):
            row_fill = fill_l1 if level == 1 else (fill_l2 if level == 2 else fill_l3)
            row_font = font_l1 if level == 1 else (font_l2 if level == 2 else font_l3)
            for col_idx, val in enumerate(r_vals, start=1):
                cell = ws.cell(row=r_offset, column=col_idx, value=val)
                cell.fill = row_fill
                cell.border = border_all
                if col_idx in (5, 6) and isinstance(val, str) and val.startswith("http"):
                    cell.hyperlink = val
                    cell.font = font_link
                else:
                    cell.font = row_font
                cell.alignment = align_center if col_idx in (1, 8) else align_left
            ws.row_dimensions[r_offset].height = 24 if level in (1, 2) else 21

        col_widths = [20, 34, 26, 28, 46, 44, 24, 22, 32]
        for idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = "A2"

    def write_to_google_sheet(
        self,
        google_client: GoogleCloudClient,
        batch_result: Dict[str, Any],
        sheet_id_or_url: str = "",
        new_sheet_title: str = "",
        merge_hierarchy_cells: bool = True,
    ) -> Tuple[bool, str, str]:
        """
        将层级上传数据写入 Google Sheets（谷歌在线表格）。
        - 如果 sheet_id_or_url 不为空：打开该表格并追加/写入
        - 如果 sheet_id_or_url 为空：自动新建一个 Google Sheet，开启公开查看权限，并写入两个层级工作表
        返回: (success, sheet_url, message)
        """
        if not google_client or not getattr(google_client, "gspread_client", None):
            return False, "", "未启用真实 Google Sheets 客户端（当前为模拟模式或未登录）"

        try:
            gc = google_client.gspread_client
            target_id = extract_google_id(sheet_id_or_url)
            created_new = False

            if target_id:
                self._log(f"正在打开指定的谷歌表格 (ID: {target_id})...", "INFO")
                sh = gc.open_by_key(target_id)
            else:
                if not new_sheet_title:
                    master_name = batch_result.get("master_folder", {}).get(
                        "name", "云端上传汇总"
                    )
                    new_sheet_title = f"【云端上传层级表】{master_name}"
                self._log(f"正在云端新建谷歌表格: {new_sheet_title} ...", "INFO")
                sh = gc.create(new_sheet_title)
                created_new = True
                # 为新建的谷歌表格开启公开查看权限，方便直接打开链接查看
                google_client.set_public_permission(sh.id)

            sheet_url = f"https://docs.google.com/spreadsheets/d/{sh.id}/edit"

            # 1. 准备「层级分栏总表」数据与合并请求
            self._write_gsheet_column_hierarchy(
                sh, batch_result, created_new, merge_hierarchy_cells
            )

            # 2. 准备「树状目录层级表」数据与高亮配色
            self._write_gsheet_tree_hierarchy(sh, batch_result)

            self._log(f"谷歌表格层级数据写入完成！链接: {sheet_url}", "SUCCESS")
            return True, sheet_url, "谷歌表格写入成功"

        except Exception as e:
            err = f"写入谷歌表格失败: {str(e)}"
            self._log(err, "ERROR")
            return False, "", err

    def _write_gsheet_column_hierarchy(
        self,
        sh,
        batch_result: Dict[str, Any],
        is_new_spreadsheet: bool,
        merge_hierarchy_cells: bool,
    ):
        """在 Google Sheet 中写入或追加「层级分栏总表」并执行单元格层级合并与上色"""
        tab_name = "层级分栏总表"
        try:
            ws = sh.worksheet(tab_name)
            existing_rows = len(ws.get_all_values())
        except gspread.WorksheetNotFound:
            if is_new_spreadsheet and len(sh.worksheets()) == 1:
                ws = sh.sheet1
                ws.update_title(tab_name)
                existing_rows = 0
            else:
                ws = sh.add_worksheet(title=tab_name, rows=500, cols=16)
                existing_rows = 0

        rows_to_write = []
        merges_to_apply = []  # (start_row_0idx, end_row_0idx, start_col_0idx, end_col_0idx)

        if existing_rows == 0:
            # 写入双层表头
            group_header = [
                "【第一层级】云端总文件夹 (Master Folder)",
                "",
                "",
                "【第二层级】分类子文件夹 (Sub-Folder)",
                "",
                "",
                "",
                "【第三层级】夹内单个文件 (Individual File & 下载链接)",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
            col_header = [
                "总文件夹名称",
                "总文件夹 ID",
                "总文件夹链接",
                "子文件夹序号",
                "子文件夹名称",
                "子文件夹 ID",
                "子文件夹链接",
                "夹内序号",
                "单个文件名称",
                "单个文件 ID",
                "单文件直接下载链接 (Download直链)",
                "单文件在线预览链接 (View链接)",
                "文件大小",
                "权限状态 / 上传时间",
            ]
            rows_to_write.append(group_header)
            rows_to_write.append(col_header)
            # 合并第一行大类表头
            merges_to_apply.append((0, 1, 0, 3))
            merges_to_apply.append((0, 1, 3, 7))
            merges_to_apply.append((0, 1, 7, 14))
            start_data_row_0idx = 2
        else:
            start_data_row_0idx = existing_rows

        master = batch_result.get("master_folder", {})
        master_name = master.get("name", "")
        master_id = master.get("id", "")
        master_link = master.get("webViewLink", "")
        subfolders = batch_result.get("subfolders", [])

        batch_row_start_0idx = start_data_row_0idx
        curr_row_0idx = start_data_row_0idx

        for sf_idx, sf in enumerate(subfolders):
            sf_order = sf.get("index", sf_idx + 1)
            sf_name = sf.get("name", "")
            sf_id = sf.get("id", "")
            sf_link = sf.get("webViewLink", "")
            files = sf.get("files", [])

            sf_start_0idx = curr_row_0idx
            if not files:
                rows_to_write.append(
                    [
                        master_name,
                        master_id,
                        master_link,
                        sf_order,
                        sf_name,
                        sf_id,
                        sf_link,
                        "-",
                        "(无文件)",
                        "-",
                        "-",
                        "-",
                        "-",
                        "无文件",
                    ]
                )
                curr_row_0idx += 1
                continue

            for f_item in files:
                rows_to_write.append(
                    [
                        master_name,
                        master_id,
                        master_link,
                        sf_order,
                        sf_name,
                        sf_id,
                        sf_link,
                        f_item.get("index_in_folder", 1),
                        f_item.get("name", ""),
                        f_item.get("id", ""),
                        f_item.get("downloadLink", ""),
                        f_item.get("webViewLink", ""),
                        f_item.get("size_str", ""),
                        f"{f_item.get('permission_status', '已开启公开权限')} ({f_item.get('upload_time', '')})",
                    ]
                )
                curr_row_0idx += 1

            sf_end_0idx = curr_row_0idx
            if merge_hierarchy_cells and (sf_end_0idx - sf_start_0idx) > 1:
                for c_idx in range(3, 7):
                    merges_to_apply.append((sf_start_0idx, sf_end_0idx, c_idx, c_idx + 1))

        batch_row_end_0idx = curr_row_0idx
        if merge_hierarchy_cells and (batch_row_end_0idx - batch_row_start_0idx) > 1:
            for c_idx in range(0, 3):
                merges_to_apply.append(
                    (batch_row_start_0idx, batch_row_end_0idx, c_idx, c_idx + 1)
                )

        if rows_to_write:
            ws.append_rows(rows_to_write, value_input_option="USER_ENTERED")

        # 构建 Google Sheets batchUpdate 格式与合并请求
        sheet_id = ws.id
        requests_body = []

        for s_r, e_r, s_c, e_c in merges_to_apply:
            requests_body.append(
                {
                    "mergeCells": {
                        "range": {
                            "sheetId": sheet_id,
                            "startRowIndex": s_r,
                            "endRowIndex": e_r,
                            "startColumnIndex": s_c,
                            "endColumnIndex": e_c,
                        },
                        "mergeType": "MERGE_ALL",
                    }
                }
            )

        if existing_rows == 0:
            # 冻结前两行表头并设置背景色
            requests_body.append(
                {
                    "updateSheetProperties": {
                        "properties": {
                            "sheetId": sheet_id,
                            "gridProperties": {"frozenRowCount": 2},
                        },
                        "fields": "gridProperties.frozenRowCount",
                    }
                }
            )
            # 第一行表头颜色
            header_colors = [
                (0, 3, {"red": 0.11, "green": 0.31, "blue": 0.45}),
                (3, 7, {"red": 0.07, "green": 0.47, "blue": 0.39}),
                (7, 14, {"red": 0.18, "green": 0.25, "blue": 0.33}),
            ]
            for sc, ec, bg in header_colors:
                requests_body.append(
                    {
                        "repeatCell": {
                            "range": {
                                "sheetId": sheet_id,
                                "startRowIndex": 0,
                                "endRowIndex": 1,
                                "startColumnIndex": sc,
                                "endColumnIndex": ec,
                            },
                            "cell": {
                                "userEnteredFormat": {
                                    "backgroundColor": bg,
                                    "horizontalAlignment": "CENTER",
                                    "verticalAlignment": "MIDDLE",
                                    "textFormat": {
                                        "foregroundColor": {
                                            "red": 1.0,
                                            "green": 1.0,
                                            "blue": 1.0,
                                        },
                                        "bold": True,
                                        "fontSize": 11,
                                    },
                                }
                            },
                            "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment)",
                        }
                    }
                )

        # 垂直居中对齐数据区
        requests_body.append(
            {
                "repeatCell": {
                    "range": {
                        "sheetId": sheet_id,
                        "startRowIndex": start_data_row_0idx,
                        "endRowIndex": batch_row_end_0idx,
                        "startColumnIndex": 0,
                        "endColumnIndex": 14,
                    },
                    "cell": {
                        "userEnteredFormat": {
                            "verticalAlignment": "MIDDLE",
                        }
                    },
                    "fields": "userEnteredFormat.verticalAlignment",
                }
            }
        )

        if requests_body:
            sh.batch_update({"requests": requests_body})

    def _write_gsheet_tree_hierarchy(self, sh, batch_result: Dict[str, Any]):
        """在 Google Sheet 中写入或追加「树状目录层级表」"""
        tab_name = "树状目录层级表"
        try:
            ws = sh.worksheet(tab_name)
            existing_rows = len(ws.get_all_values())
        except gspread.WorksheetNotFound:
            ws = sh.add_worksheet(title=tab_name, rows=500, cols=10)
            existing_rows = 0

        rows_to_write = []
        row_levels = []  # (row_0idx, level)

        if existing_rows == 0:
            rows_to_write.append(
                [
                    "层级级别",
                    "层级树状目录展示",
                    "节点名称 (文件夹名/文件名)",
                    "云端 ID (Folder / File ID)",
                    "单文件直接下载链接 (Download直链)",
                    "云端打开/预览链接 (总夹/子夹/单文件链接)",
                    "所属上级路径",
                    "文件大小 / 数量统计",
                    "权限与时间说明",
                ]
            )
            curr_0idx = 1
        else:
            curr_0idx = existing_rows

        master = batch_result.get("master_folder", {})
        master_name = master.get("name", "")
        master_id = master.get("id", "")
        master_link = master.get("webViewLink", "")
        subfolders = batch_result.get("subfolders", [])
        total_files = batch_result.get("total_uploaded_files", 0)

        # Level 1: 总文件夹
        rows_to_write.append(
            [
                "第1级：总文件夹",
                f"📁 [总文件夹] {master_name}",
                master_name,
                master_id,
                "(总文件夹请使用右侧打开链接)",
                master_link,
                "云端根目录",
                f"共 {len(subfolders)} 个子文件夹 / {total_files} 个文件",
                f"已开启公开访问权限 | {batch_result.get('batch_time', '')}",
            ]
        )
        row_levels.append((curr_0idx, 1))
        curr_0idx += 1

        # Level 2 & 3: 子文件夹与单个文件
        for sf in subfolders:
            sf_idx = sf.get("index", 1)
            sf_name = sf.get("name", "")
            sf_id = sf.get("id", "")
            sf_link = sf.get("webViewLink", "")
            files = sf.get("files", [])

            rows_to_write.append(
                [
                    f"第2级：子文件夹 #{sf_idx}",
                    f"    📂 [{sf_idx:02d}] {sf_name}",
                    sf_name,
                    sf_id,
                    "(子文件夹请使用右侧打开链接)",
                    sf_link,
                    master_name,
                    f"本夹内共 {len(files)} 个文件",
                    "已开启公开访问权限",
                ]
            )
            row_levels.append((curr_0idx, 2))
            curr_0idx += 1

            for f_item in files:
                f_idx = f_item.get("index_in_folder", 1)
                f_name = f_item.get("name", "")
                f_id = f_item.get("id", "")
                rows_to_write.append(
                    [
                        f"第3级：单个文件 ({f_idx}/{len(files)})",
                        f"        📄 ({f_idx}) {f_name}",
                        f_name,
                        f_id,
                        f_item.get("downloadLink", ""),
                        f_item.get("webViewLink", ""),
                        f"{master_name} / {sf_name}",
                        f_item.get("size_str", ""),
                        f"{f_item.get('permission_status', '')} | {f_item.get('upload_time', '')}",
                    ]
                )
                row_levels.append((curr_0idx, 3))
                curr_0idx += 1

        if rows_to_write:
            ws.append_rows(rows_to_write, value_input_option="USER_ENTERED")

        # 为树状目录的 第1级（总文件夹）和 第2级（子文件夹）行上色区分层级
        sheet_id = ws.id
        requests_body = []
        if existing_rows == 0:
            requests_body.append(
                {
                    "updateSheetProperties": {
                        "properties": {
                            "sheetId": sheet_id,
                            "gridProperties": {"frozenRowCount": 1},
                        },
                        "fields": "gridProperties.frozenRowCount",
                    }
                }
            )
            requests_body.append(
                {
                    "repeatCell": {
                        "range": {
                            "sheetId": sheet_id,
                            "startRowIndex": 0,
                            "endRowIndex": 1,
                            "startColumnIndex": 0,
                            "endColumnIndex": 9,
                        },
                        "cell": {
                            "userEnteredFormat": {
                                "backgroundColor": {
                                    "red": 0.11,
                                    "green": 0.16,
                                    "blue": 0.20,
                                },
                                "textFormat": {
                                    "foregroundColor": {
                                        "red": 1.0,
                                        "green": 1.0,
                                        "blue": 1.0,
                                    },
                                    "bold": True,
                                },
                            }
                        },
                        "fields": "userEnteredFormat(backgroundColor,textFormat)",
                    }
                }
            )

        for r_idx, lvl in row_levels:
            if lvl == 1:
                bg = {"red": 0.84, "green": 0.92, "blue": 0.97}
                bold = True
            elif lvl == 2:
                bg = {"red": 0.82, "green": 0.95, "blue": 0.92}
                bold = True
            else:
                continue
            requests_body.append(
                {
                    "repeatCell": {
                        "range": {
                            "sheetId": sheet_id,
                            "startRowIndex": r_idx,
                            "endRowIndex": r_idx + 1,
                            "startColumnIndex": 0,
                            "endColumnIndex": 9,
                        },
                        "cell": {
                            "userEnteredFormat": {
                                "backgroundColor": bg,
                                "textFormat": {"bold": bold},
                            }
                        },
                        "fields": "userEnteredFormat(backgroundColor,textFormat.bold)",
                    }
                }
            )

        if requests_body:
            sh.batch_update({"requests": requests_body})
