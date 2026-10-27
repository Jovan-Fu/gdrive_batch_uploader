import datetime
from pathlib import Path
from typing import Callable, Dict, Any, List, Optional

import openpyxl
from PyQt6.QtCore import Qt, QTimer
from PyQt6.QtGui import QFont, QTextCursor, QColor, QGuiApplication
from PyQt6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QGroupBox,
    QLabel,
    QPushButton,
    QComboBox,
    QCheckBox,
    QTextEdit,
    QTableWidget,
    QTableWidgetItem,
    QHeaderView,
    QProgressBar,
    QFileDialog,
    QMessageBox,
    QSplitter,
)

from core.google_client import parse_drive_targets_from_text
from core.deleter_worker import BatchDeleteWorker


class BatchDeletePage(QWidget):
    """
    独立页面（Tab 2）：云端链接批量删除清理页
    用于服务结束/下载完成后，将之前生成的单文件下载链接、预览链接、子文件夹链接、总文件夹链接或纯ID
    批量粘贴进来，一键执行彻底删除、移入回收站或关闭下载权限。
    """

    def __init__(
        self,
        get_base_config_cb: Callable[[], Dict[str, Any]],
        parent=None,
    ):
        super().__init__(parent)
        self.get_base_config_cb = get_base_config_cb
        self.delete_worker: Optional[BatchDeleteWorker] = None
        self.parsed_targets: List[Dict[str, str]] = []

        self.parse_timer = QTimer(self)
        self.parse_timer.setSingleShot(True)
        self.parse_timer.setInterval(250)
        self.parse_timer.timeout.connect(self._do_parse_input)

        self._init_ui()

    def _init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(8)

        grp_top = QGroupBox(
            "① 粘贴或导入待删除的谷歌云盘链接 / 文件ID（支持从谷歌表格或Excel整列/整表直接复制粘贴，自动提取ID并去重）"
        )
        top_v = QVBoxLayout(grp_top)
        top_v.setSpacing(6)

        btn_bar = QHBoxLayout()

        btn_paste = QPushButton("📋 从剪贴板一键追加粘贴")
        btn_paste.setStyleSheet(
            "background-color:#EBF5FB; color:#1B4F72; font-weight:bold;"
        )
        btn_paste.clicked.connect(self._on_paste_clipboard)
        btn_bar.addWidget(btn_paste)

        btn_import_excel = QPushButton("📊 从本地 Excel 层级报表导入链接...")
        btn_import_excel.setStyleSheet(
            "background-color:#E8F8F5; color:#0E6251; font-weight:bold;"
        )
        btn_import_excel.clicked.connect(self._on_import_from_excel)
        btn_bar.addWidget(btn_import_excel)

        btn_parse_now = QPushButton("🔍 立即解析去重并刷新右侧列表")
        btn_parse_now.clicked.connect(self._do_parse_input)
        btn_bar.addWidget(btn_parse_now)

        btn_clear_input = QPushButton("🧹 清空输入框与列表")
        btn_clear_input.clicked.connect(self._on_clear_all)
        btn_bar.addWidget(btn_clear_input)

        btn_bar.addStretch()

        self.lbl_parsed_count = QLabel("已识别去重目标: 0 个")
        self.lbl_parsed_count.setFixedHeight(26)
        self.lbl_parsed_count.setStyleSheet(
            "background-color:#FDEDEC; color:#922B21; padding:3px 10px; border-radius:4px; font-weight:bold;"
        )
        btn_bar.addWidget(self.lbl_parsed_count)

        top_v.addLayout(btn_bar)

        mid_splitter = QSplitter(Qt.Orientation.Horizontal)

        left_box = QWidget()
        left_v = QVBoxLayout(left_box)
        left_v.setContentsMargins(0, 0, 0, 0)
        left_v.addWidget(
            QLabel(
                "👇 请在此粘贴要删除的下载链接、查看链接、文件夹链接或文件ID（每行一个，或直接粘贴表格多列内容）："
            )
        )
        self.txt_links_input = QTextEdit()
        self.txt_links_input.setPlaceholderText(
            "示例（支持混合粘贴，程序会自动提取其中的云端 ID 并自动去重）：\n"
            "https://drive.google.com/uc?export=download&id=1abcDEF...\n"
            "https://drive.google.com/file/d/1abcDEF.../view?usp=sharing\n"
            "https://drive.google.com/drive/folders/1xyzUVW...?usp=sharing\n"
            "或者直接从谷歌表格/Excel复制整行整列粘贴到这里..."
        )
        self.txt_links_input.setFont(QFont("Consolas", 10))
        self.txt_links_input.textChanged.connect(
            lambda: self.parse_timer.start()
        )
        left_v.addWidget(self.txt_links_input, stretch=1)
        mid_splitter.addWidget(left_box)

        right_box = QWidget()
        right_v = QVBoxLayout(right_box)
        right_v.setContentsMargins(0, 0, 0, 0)
        right_v.addWidget(
            QLabel("📋 自动识别出的去重云端目标清单 & 实时清理状态：")
        )

        self.table_targets = QTableWidget(0, 5)
        self.table_targets.setHorizontalHeaderLabels(
            ["序号", "识别类型", "提取的云端 ID", "执行状态", "云端名称 / 结果说明"]
        )
        self.table_targets.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.ResizeMode.ResizeToContents
        )
        self.table_targets.horizontalHeader().setSectionResizeMode(
            1, QHeaderView.ResizeMode.ResizeToContents
        )
        self.table_targets.horizontalHeader().setSectionResizeMode(
            2, QHeaderView.ResizeMode.Interactive
        )
        self.table_targets.setColumnWidth(2, 210)
        self.table_targets.horizontalHeader().setSectionResizeMode(
            3, QHeaderView.ResizeMode.ResizeToContents
        )
        self.table_targets.horizontalHeader().setSectionResizeMode(
            4, QHeaderView.ResizeMode.Stretch
        )
        self.table_targets.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.table_targets.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.table_targets.setAlternatingRowColors(True)
        right_v.addWidget(self.table_targets, stretch=1)
        mid_splitter.addWidget(right_box)

        mid_splitter.setSizes([450, 620])
        top_v.addWidget(mid_splitter, stretch=1)

        grp_opts = QGroupBox("② 批量删除 / 清理方式设置")
        opts_h = QHBoxLayout(grp_opts)
        opts_h.setSpacing(12)

        opts_h.addWidget(QLabel("清理方式:"))
        self.combo_delete_mode = QComboBox()
        self.combo_delete_mode.addItem(
            "🔥 彻底永久删除（从谷歌云盘直接抹除，立即释放空间，不可恢复）",
            "permanent",
        )
        self.combo_delete_mode.addItem(
            "♻️ 移入云盘回收站（外链停止下载，30天内可在云盘回收站还原）",
            "trash",
        )
        self.combo_delete_mode.addItem(
            "🔒 仅关闭公开下载权限（文件保留在云盘内，但分享下载直链立即失效）",
            "revoke_public",
        )
        self.combo_delete_mode.setMinimumWidth(420)
        opts_h.addWidget(self.combo_delete_mode)

        self.chk_fetch_name = QCheckBox(
            "删除前查询云端文件名（日志显示具体文件名）"
        )
        self.chk_fetch_name.setChecked(True)
        opts_h.addWidget(self.chk_fetch_name)

        self.chk_auto_remove_done = QCheckBox(
            "清理完成后自动从左侧输入框剔除已成功删除的链接"
        )
        self.chk_auto_remove_done.setChecked(True)
        opts_h.addWidget(self.chk_auto_remove_done)

        opts_h.addStretch()

        main_v_splitter = QSplitter(Qt.Orientation.Vertical)
        top_container = QWidget()
        tc_layout = QVBoxLayout(top_container)
        tc_layout.setContentsMargins(0, 0, 0, 0)
        tc_layout.addWidget(grp_top, stretch=1)
        tc_layout.addWidget(grp_opts)
        main_v_splitter.addWidget(top_container)

        bottom_box = QWidget()
        bottom_v = QVBoxLayout(bottom_box)
        bottom_v.setContentsMargins(0, 4, 0, 0)
        bottom_v.setSpacing(5)

        ctrl_h = QHBoxLayout()
        self.btn_start_delete = QPushButton(
            "🗑️ 立即开始批量删除 / 清理云端链接"
        )
        self.btn_start_delete.setFixedHeight(38)
        self.btn_start_delete.setFont(
            QFont("Microsoft YaHei", 11, QFont.Weight.Bold)
        )
        self.btn_start_delete.setStyleSheet(
            """
            QPushButton {
                background-color: #C0392B;
                color: #FFFFFF;
                border: none;
                border-radius: 5px;
                padding: 6px 22px;
            }
            QPushButton:hover {
                background-color: #E74C3C;
            }
            QPushButton:disabled {
                background-color: #95A5A6;
            }
            """
        )
        self.btn_start_delete.clicked.connect(self._on_start_batch_delete)
        ctrl_h.addWidget(self.btn_start_delete)

        self.btn_stop_delete = QPushButton("🛑 停止删除")
        self.btn_stop_delete.setFixedHeight(38)
        self.btn_stop_delete.setEnabled(False)
        self.btn_stop_delete.setStyleSheet(
            """
            QPushButton {
                background-color: #7F8C8D;
                color: #FFFFFF;
                border: none;
                border-radius: 5px;
                padding: 6px 16px;
                font-weight: bold;
            }
            QPushButton:disabled {
                background-color: #BDC3C7;
            }
            """
        )
        self.btn_stop_delete.clicked.connect(self._on_stop_batch_delete)
        ctrl_h.addWidget(self.btn_stop_delete)

        self.progress_bar = QProgressBar()
        self.progress_bar.setFixedHeight(26)
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        self.progress_bar.setAlignment(Qt.AlignmentFlag.AlignCenter)
        ctrl_h.addWidget(self.progress_bar, stretch=1)

        btn_clear_log = QPushButton("🧹 清空删除日志")
        btn_clear_log.clicked.connect(lambda: self.txt_delete_log.clear())
        ctrl_h.addWidget(btn_clear_log)

        bottom_v.addLayout(ctrl_h)

        self.lbl_status = QLabel("就绪：请粘贴待删除的云端链接或 ID...")
        self.lbl_status.setFixedHeight(22)
        self.lbl_status.setWordWrap(False)
        self.lbl_status.setStyleSheet("color: #34495E; font-weight: bold;")
        bottom_v.addWidget(self.lbl_status)

        self.txt_delete_log = QTextEdit()
        self.txt_delete_log.setReadOnly(True)
        self.txt_delete_log.setFont(QFont("Consolas", 10))
        self.txt_delete_log.setStyleSheet(
            "background-color: #1E272E; color: #D2DAE2; border-radius: 6px; padding: 6px;"
        )
        bottom_v.addWidget(self.txt_delete_log, stretch=1)

        main_v_splitter.addWidget(bottom_box)
        main_v_splitter.setSizes([490, 260])

        layout.addWidget(main_v_splitter, stretch=1)

        self._append_log(
            "🗑️ 欢迎使用【云端链接批量删除清理】专页！可直接从谷歌表格或本地 Excel 报表复制下载链接、预览链接、文件夹链接或文件 ID 粘贴到上方输入框进行一键批量清理。",
            "INFO",
        )

    def _set_status_single_line(self, text: str, is_error: bool = False):
        first_line = (text or "").splitlines()[0] if text else ""
        clean = " ".join(first_line.split())
        if len(clean) > 110:
            clean = clean[:110] + "..."
        self.lbl_status.setText(clean)
        if is_error:
            self.lbl_status.setStyleSheet("color: #C0392B; font-weight: bold;")
        else:
            self.lbl_status.setStyleSheet("color: #34495E; font-weight: bold;")

    def _append_log(self, msg: str, level: str = "INFO"):
        now_str = datetime.datetime.now().strftime("%H:%M:%S")
        color_map = {
            "INFO": "#D2DAE2",
            "SUCCESS": "#2ECC71",
            "WARN": "#F39C12",
            "ERROR": "#FF5E57",
        }
        color = color_map.get(level.upper(), "#D2DAE2")
        for line in (msg or "").splitlines():
            escaped = (
                line.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
            html_line = f'<div style="color:{color}; margin-bottom:2px;">[{now_str}] [{level}] {escaped}</div>'
            self.txt_delete_log.moveCursor(QTextCursor.MoveOperation.End)
            self.txt_delete_log.insertHtml(html_line)
            self.txt_delete_log.insertPlainText("\n")
        self.txt_delete_log.moveCursor(QTextCursor.MoveOperation.End)

    def _on_paste_clipboard(self):
        cb = QGuiApplication.clipboard()
        text = cb.text() if cb else ""
        if not text or not text.strip():
            QMessageBox.information(self, "提示", "当前剪贴板中没有文本内容！")
            return
        existing = self.txt_links_input.toPlainText().strip()
        if existing:
            self.txt_links_input.setPlainText(existing + "\n" + text.strip())
        else:
            self.txt_links_input.setPlainText(text.strip())
        self._do_parse_input()

    def _on_import_from_excel(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择要导入链接的本地 Excel 层级报表",
            "",
            "Excel Files (*.xlsx);;All Files (*)",
        )
        if not path:
            return

        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = (
                wb["层级分栏总表"]
                if "层级分栏总表" in wb.sheetnames
                else wb.active
            )

            box = QMessageBox(self)
            box.setIcon(QMessageBox.Icon.Question)
            box.setWindowTitle("选择要从 Excel 报表导入的删除层级")
            box.setText(
                f"已读取报表：{Path(path).name}\n\n请选择您需要导入并清理的链接范围："
            )
            btn_files = box.addButton(
                "📄 仅导入所有【单文件下载链接】",
                QMessageBox.ButtonRole.AcceptRole,
            )
            btn_subfolders = box.addButton(
                "📂 导入所有【分类子文件夹链接】(删子夹会自动删内部单文件)",
                QMessageBox.ButtonRole.ActionRole,
            )
            btn_all = box.addButton(
                "📁 导入全部链接 (总文件夹 + 子文件夹 + 单文件)",
                QMessageBox.ButtonRole.ActionRole,
            )
            btn_cancel = box.addButton("取消", QMessageBox.ButtonRole.RejectRole)
            box.exec()

            clicked = box.clickedButton()
            if clicked == btn_cancel:
                return

            extracted_lines = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                for cell_val in row:
                    if not cell_val or not isinstance(cell_val, str):
                        continue
                    s = cell_val.strip()
                    if not s.startswith("http"):
                        continue
                    if clicked == btn_files:
                        if "export=download" in s or "/file/d/" in s:
                            extracted_lines.append(s)
                    elif clicked == btn_subfolders:
                        if "/folders/" in s:
                            extracted_lines.append(s)
                    else:
                        extracted_lines.append(s)

            if clicked == btn_subfolders and ws.max_column >= 7:
                sub_only = []
                for r_idx in range(3, ws.max_row + 1):
                    v = ws.cell(row=r_idx, column=7).value
                    if isinstance(v, str) and v.startswith("http"):
                        sub_only.append(v.strip())
                if sub_only:
                    extracted_lines = sub_only

            if not extracted_lines:
                QMessageBox.warning(
                    self, "未找到链接", "在该 Excel 表格中未提取到匹配的谷歌云盘链接！"
                )
                return

            existing = self.txt_links_input.toPlainText().strip()
            joined = "\n".join(extracted_lines)
            if existing:
                self.txt_links_input.setPlainText(existing + "\n" + joined)
            else:
                self.txt_links_input.setPlainText(joined)

            self._do_parse_input()
            self._append_log(
                f"📊 已从 Excel 报表「{Path(path).name}」导入链接，去重后共识别出 {len(self.parsed_targets)} 个云端目标。",
                "SUCCESS",
            )
        except Exception as e:
            QMessageBox.critical(self, "读取 Excel 失败", f"导入失败: {e}")

    def _on_clear_all(self):
        self.txt_links_input.clear()
        self.parsed_targets = []
        self.table_targets.setRowCount(0)
        self.lbl_parsed_count.setText("已识别去重目标: 0 个")
        self.lbl_parsed_count.setStyleSheet(
            "background-color:#FDEDEC; color:#922B21; padding:3px 10px; border-radius:4px; font-weight:bold;"
        )

    def _do_parse_input(self):
        raw_text = self.txt_links_input.toPlainText()
        self.parsed_targets = parse_drive_targets_from_text(raw_text)
        count = len(self.parsed_targets)

        if count == 0:
            self.lbl_parsed_count.setText("已识别去重目标: 0 个")
            self.lbl_parsed_count.setStyleSheet(
                "background-color:#FDEDEC; color:#922B21; padding:3px 10px; border-radius:4px; font-weight:bold;"
            )
        else:
            self.lbl_parsed_count.setText(f"✅ 已识别去重目标: {count} 个")
            self.lbl_parsed_count.setStyleSheet(
                "background-color:#D4EFDF; color:#196F3D; padding:3px 10px; border-radius:4px; font-weight:bold;"
            )

        self.table_targets.setRowCount(count)
        for idx, item in enumerate(self.parsed_targets):
            self.table_targets.setItem(idx, 0, QTableWidgetItem(str(idx + 1)))
            self.table_targets.setItem(
                idx, 1, QTableWidgetItem(item.get("type_hint", "链接/ID"))
            )
            self.table_targets.setItem(idx, 2, QTableWidgetItem(item["id"]))
            status_item = QTableWidgetItem("⏳ 待处理")
            status_item.setForeground(QColor("#7F8C8D"))
            self.table_targets.setItem(idx, 3, status_item)
            self.table_targets.setItem(
                idx, 4, QTableWidgetItem(item.get("raw_input", ""))
            )

    def _on_start_batch_delete(self):
        self._do_parse_input()
        if not self.parsed_targets:
            QMessageBox.warning(
                self,
                "未识别到有效链接",
                "请先在左侧输入框粘贴需要删除的 Google Drive 下载链接、预览链接、文件夹链接或文件 ID！",
            )
            return

        base_cfg = self.get_base_config_cb()
        cred = base_cfg.get("cred_file", "")
        if not cred or not Path(cred).is_file():
            QMessageBox.warning(
                self,
                "缺少 Google 授权凭证",
                "请先在窗口顶部的「🔑 Google 云端账号授权与网络设置」中选择您的 JSON 凭证文件！",
            )
            return

        delete_mode = self.combo_delete_mode.currentData()
        mode_text = self.combo_delete_mode.currentText()
        count = len(self.parsed_targets)

        confirm = QMessageBox.question(
            self,
            "确认执行批量删除/清理",
            f"您确定要对已识别出的 【{count}】 个云端目标执行以下操作吗？\n\n"
            f"👉 操作方式：{mode_text}\n\n"
            f"点击「Yes」立即开始执行。",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if confirm != QMessageBox.StandardButton.Yes:
            return

        worker_cfg = dict(base_cfg)
        worker_cfg["delete_mode"] = delete_mode
        worker_cfg["fetch_name_first"] = self.chk_fetch_name.isChecked()

        self.btn_start_delete.setEnabled(False)
        self.btn_stop_delete.setEnabled(True)
        self.progress_bar.setValue(0)
        self._set_status_single_line("正在启动批量删除线程...", is_error=False)

        self.delete_worker = BatchDeleteWorker(
            config=worker_cfg, targets=self.parsed_targets, parent=self
        )
        self.delete_worker.log_signal.connect(self._append_log)
        self.delete_worker.progress_signal.connect(self._on_worker_progress)
        self.delete_worker.item_result_signal.connect(self._on_item_result)
        self.delete_worker.finished_signal.connect(self._on_worker_finished)
        self.delete_worker.start()

    def _on_stop_batch_delete(self):
        if self.delete_worker and self.delete_worker.isRunning():
            self.delete_worker.request_stop()
            self._append_log("正在请求停止批量删除任务...", "WARN")
            self.btn_stop_delete.setEnabled(False)

    def _on_worker_progress(self, done: int, total: int, status_text: str):
        pct = int(done * 100 / max(total, 1))
        self.progress_bar.setValue(pct)
        self._set_status_single_line(status_text, is_error=False)

    def _on_item_result(
        self,
        row_idx: int,
        fid: str,
        item_name: str,
        status_code: str,
        detail_msg: str,
    ):
        if row_idx < 0 or row_idx >= self.table_targets.rowCount():
            return

        if status_code == "SUCCESS":
            st_item = QTableWidgetItem("✅ 已清理")
            st_item.setForeground(QColor("#196F3D"))
        elif status_code == "NOT_FOUND":
            st_item = QTableWidgetItem("⚪ 已不存在(跳过)")
            st_item.setForeground(QColor("#B9770E"))
        else:
            st_item = QTableWidgetItem("❌ 失败")
            st_item.setForeground(QColor("#C0392B"))

        self.table_targets.setItem(row_idx, 3, st_item)
        self.table_targets.setItem(row_idx, 4, QTableWidgetItem(detail_msg))
        self.table_targets.scrollToItem(st_item)

    def _on_worker_finished(self, ok: bool, summary: dict):
        self.btn_start_delete.setEnabled(True)
        self.btn_stop_delete.setEnabled(False)

        if not ok:
            err = summary.get("error", "未知错误")
            self._set_status_single_line(f"❌ 批量删除未完成: {err}", is_error=True)
            QMessageBox.warning(self, "批量删除未完成", err)
            return

        self.progress_bar.setValue(100)
        total = summary.get("total", 0)
        succ = summary.get("success_count", 0)
        nf = summary.get("not_found_count", 0)
        err_c = summary.get("error_count", 0)
        succeeded_ids = set(summary.get("succeeded_ids", []))

        self._set_status_single_line(
            f"✅ 清理完毕！共 {total} 个：成功清理 {succ} 个，已不存在(跳过) {nf} 个，失败 {err_c} 个。",
            is_error=(err_c > 0),
        )

        if self.chk_auto_remove_done.isChecked() and succeeded_ids:
            remaining_lines = []
            for line in self.txt_links_input.toPlainText().splitlines():
                line_targets = parse_drive_targets_from_text(line)
                if not line_targets:
                    continue
                if all(t["id"] in succeeded_ids for t in line_targets):
                    continue
                remaining_lines.append(line)
            self.txt_links_input.blockSignals(True)
            self.txt_links_input.setPlainText("\n".join(remaining_lines))
            self.txt_links_input.blockSignals(False)

        QMessageBox.information(
            self,
            "批量删除/清理完成",
            f"✅ 云端链接批量清理已完成！\n\n"
            f"• 执行模式：{summary.get('mode_desc', '')}\n"
            f"• 识别目标总数：{total} 个\n"
            f"• 成功清理数量：{succ} 个\n"
            f"• 云端已不存在(自动跳过)：{nf} 个\n"
            f"• 失败数量：{err_c} 个",
        )
